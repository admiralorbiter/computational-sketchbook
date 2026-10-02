"""
src/harvest_mo_assessment.py

Extracts Missouri MAP assessment performance, MPI, and participation metrics
for the 36 Missouri balanced panel districts across 2018-19 through 2023-24.

Sources:
- Missouri DESE APR Supporting Data Reports (2019, 2022, 2023, 2024)
- PRiME Center at St. Louis University Master Dataset (longitudinal MAP results)
"""

import os
import openpyxl
import pandas as pd

MO_DISTRICTS = {
    "2903200": {"mo_id": "019139", "name": "ARCHIE R-V"},
    "2904620": {"mo_id": "019152", "name": "BELTON 124"},
    "2905310": {"mo_id": "048068", "name": "BLUE SPRINGS R-IV"},
    "2908250": {"mo_id": "048080", "name": "CENTER 58"},
    "2911070": {"mo_id": "019150", "name": "DREXEL R-IV"},
    "2911160": {"mo_id": "019147", "name": "EAST LYNNE 40"},
    "2911650": {"mo_id": "024089", "name": "EXCELSIOR SPRINGS 40"},
    "2912290": {"mo_id": "048066", "name": "FORT OSAGE R-I"},
    "2913080": {"mo_id": "048069", "name": "GRAIN VALLEY R-V"},
    "2913140": {"mo_id": "048074", "name": "GRANDVIEW C-4"},
    "2913680": {"mo_id": "089088", "name": "HARDIN-CENTRAL C-2"},
    "2913760": {"mo_id": "019149", "name": "HARRISONVILLE R-IX"},
    "2914340": {"mo_id": "048072", "name": "HICKMAN MILLS C-1"},
    "2915480": {"mo_id": "048077", "name": "INDEPENDENCE 30"},
    "2916400": {"mo_id": "048078", "name": "KANSAS CITY 33"},
    "2916450": {"mo_id": "024086", "name": "KEARNEY R-I"},
    "2918220": {"mo_id": "089080", "name": "LAWSON R-XIV"},
    "2918300": {"mo_id": "048071", "name": "LEE'S SUMMIT R-VII"},
    "2918540": {"mo_id": "024090", "name": "LIBERTY 53"},
    "2919230": {"mo_id": "048075", "name": "LONE JACK C-6"},
    "2931800": {"mo_id": "019151", "name": "MIDWAY R-I"},
    "2921060": {"mo_id": "024091", "name": "MISSOURI CITY 56"},
    "2922800": {"mo_id": "024093", "name": "NORTH KANSAS CITY 74"},
    "2922830": {"mo_id": "083001", "name": "NORTH PLATTE CO. R-I"},
    "2923010": {"mo_id": "048070", "name": "OAK GROVE R-VI"},
    "2923220": {"mo_id": "089087", "name": "ORRICK R-XI"},
    "2923550": {"mo_id": "083005", "name": "PARK HILL"},
    "2925230": {"mo_id": "083003", "name": "PLATTE CO. R-III"},
    "2925330": {"mo_id": "019148", "name": "PLEASANT HILL R-III"},
    "2923730": {"mo_id": "019142", "name": "RAYMORE-PECULIAR R-II"},
    "2926070": {"mo_id": "048073", "name": "RAYTOWN C-2"},
    "2926480": {"mo_id": "089089", "name": "RICHMOND R-XVI"},
    "2910320": {"mo_id": "019144", "name": "SHERWOOD CASS R-VIII"},
    "2928410": {"mo_id": "024087", "name": "SMITHVILLE R-II"},
    "2929670": {"mo_id": "019140", "name": "STRASBURG C-3"},
    "2931710": {"mo_id": "083002", "name": "WEST PLATTE CO. R-II"},
}


def load_prime_data():
    """Extract MAP proficiency for Missouri panel districts from PRiME master dataset."""
    wb = openpyxl.load_workbook("data/raw/prime_master_sample.xlsx", read_only=True)
    ws = wb["LEA-Level Data"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)

    nces_idx = header.index("NCES ID#")
    
    # Map years and subjects to column indexes
    # Note trailing spaces or variations in PRiME column headers
    col_map = {}
    for i, col in enumerate(header):
        c = str(col).strip()
        if "MAP % Proficient/Advanced" in c:
            # Check subgroup (Total vs FRL)
            subgroup = "Total" if "(Total," in c else ("FRL" if "(FRL," in c else None)
            if not subgroup:
                continue
            subj = "ELA" if "ELA" in c else ("Math" if "Math" in c else None)
            if not subj:
                continue
            # Year extraction
            for yr_str, yr_int in [("2018-19", 2019), ("2020-21", 2021), ("2021-22", 2022), ("2022-23", 2023), ("2023-24", 2024)]:
                if yr_str in c:
                    col_map[(yr_int, subj, subgroup)] = i

    records = {}
    for r in rows:
        nid = str(r[nces_idx]).strip()
        if nid in MO_DISTRICTS:
            for (yr_int, subj, subgroup), c_idx in col_map.items():
                val = r[c_idx]
                # Preserve suppression indicator '*' as null / string indicator
                key = (nid, yr_int, subj)
                if key not in records:
                    records[key] = {}
                records[key][f"pct_prof_{subgroup.lower()}"] = val

    return records


def load_mo_apr_data():
    """Extract MPI and participant counts from official DESE APR supporting reports."""
    apr_data = {}
    
    # 2019 MSIP5
    wb19 = openpyxl.load_workbook("data/raw/mo_assessment/mo_apr_supporting_2019.xlsx", read_only=True)
    ws19 = wb19.active
    rows19 = ws19.iter_rows(values_only=True)
    h19 = next(rows19)
    code_idx = 1
    ela_mpi_idx = h19.index("S1_ELA_CURR_MPI")
    math_mpi_idx = h19.index("S1_MA_CURR_MPI")
    ela_prof_idx = h19.index("S1_ELA_CURR_PERCENT_PROF_OR_ADVANCED")
    math_prof_idx = h19.index("S1_MA_CURR_PERCENT_PROF_OR_ADVANCED")

    # Map code to NCES
    code_to_nces = {v["mo_id"]: k for k, v in MO_DISTRICTS.items()}

    for r in rows19:
        code = str(r[code_idx]).zfill(6)
        if code in code_to_nces:
            nid = code_to_nces[code]
            apr_data[(nid, 2019, "ELA")] = {
                "mpi": r[ela_mpi_idx],
                "apr_pct_prof": r[ela_prof_idx],
                "accountable_n": None,
                "tested_n": None,
                "participation_rate": None,
            }
            apr_data[(nid, 2019, "Math")] = {
                "mpi": r[math_mpi_idx],
                "apr_pct_prof": r[math_prof_idx],
                "accountable_n": None,
                "tested_n": None,
                "participation_rate": None,
            }

    # 2022 MSIP6
    wb22 = openpyxl.load_workbook("data/raw/mo_assessment/mo_apr_supporting_2022.xlsx", read_only=True)
    ws22 = wb22.active
    rows22 = ws22.iter_rows(values_only=True)
    h22 = next(rows22)
    ela_mpi_22 = h22.index("ALL_ELA_CURR_MPI")
    math_mpi_22 = h22.index("ALL_MATH_CURR_MPI")
    ela_prof_22 = h22.index("ALL_ELA_CURR_PERCENT_PROF_OR_ADVANCED")
    math_prof_22 = h22.index("ALL_MATH_CURR_PERCENT_PROF_OR_ADVANCED")

    for r in rows22:
        code = str(r[code_idx]).zfill(6)
        if code in code_to_nces:
            nid = code_to_nces[code]
            apr_data[(nid, 2022, "ELA")] = {
                "mpi": r[ela_mpi_22],
                "apr_pct_prof": r[ela_prof_22],
                "accountable_n": None,
                "tested_n": None,
                "participation_rate": None,
            }
            apr_data[(nid, 2022, "Math")] = {
                "mpi": r[math_mpi_22],
                "apr_pct_prof": r[math_prof_22],
                "accountable_n": None,
                "tested_n": None,
                "participation_rate": None,
            }

    # 2023 & 2024 MSIP6
    for yr in [2023, 2024]:
        fn = f"data/raw/mo_assessment/mo_apr_supporting_{yr}.xlsx"
        wb = openpyxl.load_workbook(fn, read_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        h = next(rows)
        
        ela_mpi = h.index("ELA_ALL_STATUS_MPI")
        math_mpi = h.index("MATH_ALL_STATUS_MPI")
        ela_acc = h.index("ELA_ALL_STATUS_ACCOUNTABLE")
        ela_part = h.index("ELA_ALL_STATUS_PARTICIPANTS")
        ela_rate = h.index("ELA_ALL_STATUS_PARTICIPATION_RATE")
        math_acc = h.index("MATH_ALL_STATUS_ACCOUNTABLE")
        math_part = h.index("MATH_ALL_STATUS_PARTICIPANTS")
        math_rate = h.index("MATH_ALL_STATUS_PARTICIPATION_RATE")

        for r in rows:
            code = str(r[code_idx]).zfill(6)
            if code in code_to_nces:
                nid = code_to_nces[code]
                apr_data[(nid, yr, "ELA")] = {
                    "mpi": r[ela_mpi],
                    "apr_pct_prof": None,
                    "accountable_n": r[ela_acc],
                    "tested_n": r[ela_part],
                    "participation_rate": r[ela_rate],
                }
                apr_data[(nid, yr, "Math")] = {
                    "mpi": r[math_mpi],
                    "apr_pct_prof": None,
                    "accountable_n": r[math_acc],
                    "tested_n": r[math_part],
                    "participation_rate": r[math_rate],
                }

    return apr_data


def main():
    os.makedirs("data/raw/mo_assessment", exist_ok=True)
    prime_records = load_prime_data()
    apr_records = load_mo_apr_data()

    rows = []
    years = [2019, 2021, 2022, 2023, 2024]
    subjects = ["ELA", "Math"]

    for nces_id, info in MO_DISTRICTS.items():
        mo_id = info["mo_id"]
        dist_name = info["name"]
        for yr in years:
            for subj in subjects:
                key = (nces_id, yr, subj)
                prime_info = prime_records.get(key, {})
                apr_info = apr_records.get(key, {})

                pct_prof_total = prime_info.get("pct_prof_total")
                pct_prof_frl = prime_info.get("pct_prof_frl")
                
                mpi = apr_info.get("mpi")
                acc_n = apr_info.get("accountable_n")
                tested_n = apr_info.get("tested_n")
                part_rate = apr_info.get("participation_rate")

                rows.append({
                    "nces_lea_id": int(nces_id),
                    "state_district_id": mo_id,
                    "district_name": dist_name,
                    "state": "MO",
                    "school_year": yr,
                    "subject": subj,
                    "pct_proficient_total": pct_prof_total,
                    "pct_proficient_frl": pct_prof_frl,
                    "mpi": mpi,
                    "accountable_n": acc_n,
                    "tested_n": tested_n,
                    "participation_rate": part_rate,
                })

    df = pd.DataFrame(rows)
    out_path = "data/raw/mo_assessment/mo_district_assessment_panel.csv"
    df.to_csv(out_path, index=False)
    print(f"Saved {len(df)} Missouri district-year-subject assessment records to {out_path}")
    print(df.head(10).to_string())


if __name__ == "__main__":
    main()
