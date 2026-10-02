"""
Parse Missouri DESE ASBR statutory Excel workbooks for focal archetypes:
- Lee's Summit R-VII (048-071)
- North Kansas City 74 (024-093)
- Raytown C-2 (048-073)
Extract Function 2200 (Instructional Support) disaggregated by Object:
- 6110 Certificated Salaries
- 6150 Non-Certificated Salaries
- 6200 Benefits
- 6300 Purchased Professional & Technical Services
- 6400 Supplies / Materials / Software
- 6500 Capital Outlay / Equipment
"""

import re
from pathlib import Path
import openpyxl
import pandas as pd

BASE_DIR = Path(__file__).resolve().parent.parent
RAW_DIR = BASE_DIR / "data" / "raw" / "missouri_asbr"
OUT_DIR = BASE_DIR / "outputs" / "tables"
OUT_DIR.mkdir(parents=True, exist_ok=True)

CPI_DEFLATORS = {
    "2014-2015": 1.2307,
    "2018-2019": 1.1578,
    "2022-2023": 1.0000,
}

# Audited enrollment from canonical panel
ENROLLMENT = {
    ("048-071", "2014-2015"): 17665,
    ("048-071", "2018-2019"): 17942,
    ("048-071", "2022-2023"): 18081,
    ("024-093", "2014-2015"): 19047,
    ("024-093", "2018-2019"): 20438,
    ("024-093", "2022-2023"): 21105,
    ("048-073", "2014-2015"): 8943,
    ("048-073", "2018-2019"): 8806,
    ("048-073", "2022-2023"): 8251,
}

CORSUP_FTE = {
    ("048-071", "2014-2015"): 21.5,
    ("048-071", "2018-2019"): 8.0,
    ("048-071", "2022-2023"): 12.3,
    ("024-093", "2014-2015"): 15.0,
    ("024-093", "2018-2019"): 21.4,
    ("024-093", "2022-2023"): 33.7,
    ("048-073", "2014-2015"): 23.2,
    ("048-073", "2018-2019"): 22.8,
    ("048-073", "2022-2023"): 22.0,
}

DISTRICTS = [
    {"slug": "lees_summit", "code": "048-071", "name": "Lee's Summit R-VII"},
    {"slug": "north_kansas_city", "code": "024-093", "name": "North Kansas City 74"},
    {"slug": "raytown", "code": "048-073", "name": "Raytown C-2"},
]

YEARS = [
    ("2014-2015", "2014_2015"),
    ("2018-2019", "2018_2019"),
    ("2022-2023", "2022_2023"),
]

def clean_num(val):
    if val is None or val == "-" or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace("$", "").replace(",", "").replace("(", "-").replace(")", "").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0

def parse_all_mo_asbr():
    records = []

    for d in DISTRICTS:
        slug = d["slug"]
        code = d["code"]
        dname = d["name"]

        for school_year, year_slug in YEARS:
            fpath = RAW_DIR / slug / f"asbr_{year_slug}.xlsx"
            if not fpath.exists():
                print(f"Missing {fpath}")
                continue

            wb = openpyxl.load_workbook(fpath, data_only=True)
            sheet = wb.active

            # Find Part III-B (Program/Object)
            obj_start_row = None
            for r in range(1, sheet.max_row + 1):
                row_text = " ".join([str(c.value or "") for c in sheet[r]])
                if "Part III-B Expenditures - Program/Object" in row_text:
                    obj_start_row = r
                    break

            if not obj_start_row:
                print(f"Could not find Part III-B in {fpath}")
                continue

            # Identify Function 2200 rows under Part III-B
            sal_cert = 0.0
            sal_noncert = 0.0
            benefits = 0.0
            purchased = 0.0
            supplies = 0.0
            capital = 0.0
            other_obj = 0.0
            total_2200 = 0.0

            fn_codes = ["2210", "2211", "2212", "2213", "2214", "2219", "2220", "2221", "2222", "2223", "2224", "2225", "2229", "2291"]

            # Scan rows after obj_start_row
            for r in range(obj_start_row, min(obj_start_row + 300, sheet.max_row + 1)):
                vals = [c.value for c in sheet[r] if c.value is not None and str(c.value).strip() != ""]
                if len(vals) >= 10 and str(vals[0]).strip() in fn_codes:
                    c_sal = clean_num(vals[2])
                    nc_sal = clean_num(vals[3])
                    ben = clean_num(vals[4])
                    pur = clean_num(vals[5])
                    sup = clean_num(vals[6])
                    cap = clean_num(vals[7])
                    oth = clean_num(vals[8])
                    tot = clean_num(vals[-1])

                    sal_cert += c_sal
                    sal_noncert += nc_sal
                    benefits += ben
                    purchased += pur
                    supplies += sup
                    capital += cap
                    other_obj += oth
                    total_2200 += tot

            enrollment = ENROLLMENT.get((code, school_year), 1)
            cpi = CPI_DEFLATORS.get(school_year, 1.0)
            corsup = CORSUP_FTE.get((code, school_year), 0.0)

            total_sal = sal_cert + sal_noncert
            total_np = purchased + supplies + capital + other_obj

            rec = {
                "district_name": dname,
                "state": "MO",
                "dese_code": code,
                "school_year": school_year,
                "enrollment": enrollment,
                "corsup_fte": corsup,
                "corsup_per_1000_pupils": round(corsup / enrollment * 1000, 3),
                # Nominal amounts
                "nom_salaries_cert_6110": sal_cert,
                "nom_salaries_noncert_6150": sal_noncert,
                "nom_salaries_total": total_sal,
                "nom_benefits_6200": benefits,
                "nom_purchased_services_6300": purchased,
                "nom_supplies_materials_6400": supplies,
                "nom_capital_outlay_6500": capital,
                "nom_other_objects_6600": other_obj,
                "nom_nonpersonnel_support_total": total_np,
                "nom_total_support_2200": total_2200,
                # Real Per-Pupil amounts (Constant 2023 Dollars)
                "real_salaries_per_pupil": round((total_sal / enrollment) * cpi, 2),
                "real_benefits_per_pupil": round((benefits / enrollment) * cpi, 2),
                "real_purchased_services_6300_per_pupil": round((purchased / enrollment) * cpi, 2),
                "real_supplies_materials_6400_per_pupil": round((supplies / enrollment) * cpi, 2),
                "real_capital_outlay_6500_per_pupil": round((capital / enrollment) * cpi, 2),
                "real_nonpersonnel_support_per_pupil": round((total_np / enrollment) * cpi, 2),
                "real_total_support_2200_per_pupil": round((total_2200 / enrollment) * cpi, 2),
                # Non-personnel component shares
                "purchased_services_share_of_np_pct": round((purchased / total_np * 100) if total_np > 0 else 0, 1),
                "supplies_materials_share_of_np_pct": round((supplies / total_np * 100) if total_np > 0 else 0, 1),
                "capital_share_of_np_pct": round((capital / total_np * 100) if total_np > 0 else 0, 1),
                "nonpersonnel_share_of_total_e07_pct": round((total_np / total_2200 * 100) if total_2200 > 0 else 0, 1),
            }
            records.append(rec)

    df_mo = pd.DataFrame(records)
    out_csv = OUT_DIR / "missouri_asbr_object_decomposition.csv"
    df_mo.to_csv(out_csv, index=False)
    print(f"Saved Missouri ASBR object decomposition to {out_csv}")
    print("\nSummary Table:")
    cols_show = [
        "district_name", "school_year", "corsup_fte",
        "real_total_support_2200_per_pupil", "real_purchased_services_6300_per_pupil",
        "real_supplies_materials_6400_per_pupil", "real_nonpersonnel_support_per_pupil",
        "purchased_services_share_of_np_pct", "supplies_materials_share_of_np_pct"
    ]
    print(df_mo[cols_show].to_string(index=False))
    return df_mo

if __name__ == "__main__":
    parse_all_mo_asbr()
